"""Format loaders: PDF, DOCX, HTML, Markdown and plain text.

Each loader returns a `LoadedDocument`: a list of `LoadedPage` (raw text per page, `number=None` when the
format has no stable page concept) plus loader warnings. Loaders only extract text; cleaning is done by
`normalize.py`. Optional libraries are imported lazily and a missing one raises `MissingDependencyError`
naming the package - there is never a fallback to a different extractor.

Page numbers are only recorded where the format defines them (PDF). DOCX/HTML/Markdown/text chunks get
`source_pages == []` rather than an invented page.
"""
from __future__ import annotations

import importlib
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Dict, List, Optional

from .errors import IngestionError, MissingDependencyError, NoTextLayerError

FORMAT_BY_SUFFIX = {
    ".pdf": "pdf", ".docx": "docx", ".html": "html", ".htm": "html",
    ".md": "markdown", ".markdown": "markdown", ".txt": "text",
}


@dataclass
class LoadedPage:
    number: Optional[int]  # 1-based page number, or None when the format is not paginated
    text: str


@dataclass
class LoadedDocument:
    format: str
    pages: List[LoadedPage]
    warnings: List[str] = field(default_factory=list)
    info: Dict[str, Any] = field(default_factory=dict)


def _require(package: str, needed_for: str, pip_name: Optional[str] = None):
    try:
        return importlib.import_module(package)
    except ImportError as exc:
        raise MissingDependencyError(package, needed_for, pip_name) from exc


def detect_format(path: Path, explicit: Optional[str] = None) -> str:
    if explicit:
        if explicit not in set(FORMAT_BY_SUFFIX.values()):
            raise IngestionError(f"unsupported format {explicit!r}; supported: {sorted(set(FORMAT_BY_SUFFIX.values()))}")
        return explicit
    fmt = FORMAT_BY_SUFFIX.get(path.suffix.lower())
    if fmt is None:
        raise IngestionError(
            f"cannot infer format of {path.name!r} from its suffix; set `format:` to one of "
            f"{sorted(set(FORMAT_BY_SUFFIX.values()))}")
    return fmt


def load_document(path: Path, fmt: Optional[str] = None, options: Optional[Dict[str, Any]] = None) -> LoadedDocument:
    path = Path(path)
    if not path.is_file():
        raise IngestionError(f"source file not found: {path}")
    fmt = detect_format(path, fmt)
    options = dict(options or {})
    loader = {"pdf": load_pdf, "docx": load_docx, "html": load_html, "markdown": load_text, "text": load_text}[fmt]
    doc = loader(path, **options) if fmt != "markdown" else loader(path, kind="markdown", **options)
    doc.format = fmt
    return doc


# ---------------------------------------------------------------------------------------------- PDF

def load_pdf(path: Path, mode: str = "layout", min_chars_per_page: int = 20) -> LoadedDocument:
    """Per-page text via pypdf. `mode="layout"` keeps word spacing correct for fonts where the default
    mode inserts spurious spaces inside words (observed on the Vietnamese Official Gazette PDF)."""
    if mode not in ("layout", "plain"):
        raise IngestionError(f"pdf mode must be 'layout' or 'plain', got {mode!r}")
    pypdf = _require("pypdf", "PDF ingestion")
    try:
        reader = pypdf.PdfReader(str(path))
    except Exception as exc:  # pypdf raises several unrelated types for corrupt files
        raise IngestionError(f"cannot open PDF {path.name}: {exc}") from exc
    if reader.is_encrypted:
        raise IngestionError(f"PDF {path.name} is encrypted; decrypt it before ingestion")
    pages: List[LoadedPage] = []
    for number, page in enumerate(reader.pages, 1):
        text = page.extract_text(extraction_mode="layout") if mode == "layout" else page.extract_text()
        pages.append(LoadedPage(number, text or ""))
    empty = [p.number for p in pages if len(p.text.strip()) < min_chars_per_page]
    if len(empty) == len(pages):
        raise NoTextLayerError(
            f"PDF {path.name} has no extractable text layer on any of its {len(pages)} pages (scanned or image-only). "
            "OCR is not implemented in src.ingestion; provide a text-layer PDF or OCR it first.")
    warnings = []
    if empty:
        warnings.append(f"{len(empty)} of {len(pages)} pages have (almost) no text layer and contribute nothing: pages {empty}")
    return LoadedDocument("pdf", pages, warnings, {"pypdf_mode": mode, "page_count": len(pages), "pages_without_text": empty})


# --------------------------------------------------------------------------------------------- DOCX

_HEADING_STYLE = re.compile(r"^heading\s*(\d)$", re.I)


def load_docx(path: Path, headings_as_markdown: bool = True) -> LoadedDocument:
    """Body paragraphs and tables of a .docx in document order. Heading styles become Markdown `#` markers when
    `headings_as_markdown` is true (a structural marker, not a text rewrite). Tables are flattened to one line
    per row with ` | ` between distinct cells. Headers, footers, footnotes and text boxes are not read."""
    docx = _require("docx", "DOCX ingestion", "python-docx")
    from docx.table import Table  # noqa: WPS433 - only importable once python-docx is present
    from docx.text.paragraph import Paragraph
    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise IngestionError(f"cannot open DOCX {path.name}: {exc}") from exc
    lines: List[str] = []
    tables = headings = 0
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name if para.style is not None else "") or ""
            match = _HEADING_STYLE.match(style)
            if headings_as_markdown and (match or style.lower() == "title"):
                level = int(match.group(1)) if match else 1
                lines.append("#" * level + " " + text)
                headings += 1
            else:
                lines.append(text)
        elif tag == "tbl":
            tables += 1
            for row in Table(child, document).rows:
                seen, cells = set(), []
                for cell in row.cells:
                    if id(cell._tc) in seen:  # merged cells repeat the same element
                        continue
                    seen.add(id(cell._tc))
                    cell_text = " ".join(cell.text.split())
                    if cell_text:
                        cells.append(cell_text)
                if cells:
                    lines.append(" | ".join(cells))
    warnings = []
    if tables:
        warnings.append(f"{tables} table(s) flattened to one line per row; cell structure is not preserved")
    return LoadedDocument("docx", [LoadedPage(None, "\n".join(lines))], warnings,
                          {"tables": tables, "style_headings": headings, "headings_as_markdown": headings_as_markdown})


# --------------------------------------------------------------------------------------------- HTML

_SKIP_TAGS = {"script", "style", "noscript", "template", "svg", "head"}
_BLOCK_TAGS = {
    "p", "div", "section", "article", "main", "header", "footer", "nav", "aside", "li", "ul", "ol", "dl", "dt", "dd",
    "tr", "table", "thead", "tbody", "tfoot", "blockquote", "pre", "figure", "figcaption", "hr", "br", "form",
    "h1", "h2", "h3", "h4", "h5", "h6",
}


class _TextExtractor(HTMLParser):
    def __init__(self, headings_as_markdown: bool, skip_tags: set):
        super().__init__(convert_charrefs=True)
        self.parts: List[str] = []
        self.skip_depth = 0
        self.skip_tags = skip_tags
        self.headings_as_markdown = headings_as_markdown
        self.title = ""
        self._in_title = False
        self.cell_open = False
        self.row_has_cell = False
        self.in_pre = 0

    def handle_starttag(self, tag, attrs):
        if tag == "title":
            self._in_title = True
        if tag in self.skip_tags:
            self.skip_depth += 1
            return
        if self.skip_depth:
            return
        if tag in ("td", "th"):
            self.parts.append(" | " if self.row_has_cell else "")
            self.row_has_cell = True
            return
        if tag == "tr":
            self.row_has_cell = False
        if tag == "pre":
            self.in_pre += 1
        if tag in _BLOCK_TAGS:
            self.parts.append("\n")
            if self.headings_as_markdown and len(tag) == 2 and tag[0] == "h" and tag[1] in "123456":
                self.parts.append("#" * int(tag[1]) + " ")

    def handle_endtag(self, tag):
        if tag == "title":
            self._in_title = False
        if tag in self.skip_tags:
            self.skip_depth = max(0, self.skip_depth - 1)
            return
        if self.skip_depth:
            return
        if tag == "pre":
            self.in_pre = max(0, self.in_pre - 1)
        if tag in _BLOCK_TAGS:
            self.parts.append("\n")

    def handle_data(self, data):
        if self._in_title:
            self.title += data
        if self.skip_depth:
            return
        self.parts.append(data if self.in_pre else re.sub(r"\s+", " ", data))


def _decode_html(raw: bytes, name: str) -> str:
    if raw.startswith(b"\xef\xbb\xbf"):
        raw = raw[3:]
    match = re.search(rb"<meta[^>]+charset\s*=\s*[\"']?\s*([A-Za-z0-9_\-]+)", raw[:4096], re.I)
    charset = match.group(1).decode("ascii") if match else "utf-8"
    try:
        return raw.decode(charset)
    except (UnicodeDecodeError, LookupError) as exc:
        raise IngestionError(f"cannot decode HTML {name} as {charset!r}: {exc}") from exc


def load_html(path: Path, headings_as_markdown: bool = True, skip_tags: Optional[List[str]] = None) -> LoadedDocument:
    """Visible text of an HTML file using the standard-library parser (no optional dependency)."""
    html = _decode_html(Path(path).read_bytes(), Path(path).name)
    parser = _TextExtractor(headings_as_markdown, set(skip_tags) if skip_tags is not None else set(_SKIP_TAGS))
    parser.feed(html)
    parser.close()
    text = "".join(parser.parts)
    return LoadedDocument("html", [LoadedPage(None, text)], [],
                          {"html_title": " ".join(parser.title.split()), "headings_as_markdown": headings_as_markdown})


# ---------------------------------------------------------------------------------- Markdown / text

def load_text(path: Path, kind: str = "text") -> LoadedDocument:
    raw = Path(path).read_bytes()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise IngestionError(f"{Path(path).name} is not valid UTF-8: {exc}") from exc
    return LoadedDocument(kind, [LoadedPage(None, text)], [], {})
